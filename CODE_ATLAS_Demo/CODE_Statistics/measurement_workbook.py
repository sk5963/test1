"""전체 가로형 측정값을 사람이 확인할 수 있는 Excel 결과로 내보낸다."""

import math
from pathlib import Path
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from .workbook import safe_cell
from CODE_Common.measurements import header, records

MAX_COLUMNS = 16384
MAX_DATA_ROWS = 1048575  # Excel 한 시트의 1048576행에서 제목 행을 제외한다.


def number(value, integer=False):
    try:
        result = float(value)
        if not math.isfinite(result):
            return None
        return int(result) if integer and result.is_integer() else result
    except (ValueError, TypeError):
        return None


def save_measurements(path, data_path, summary):
    base, tests, offset = header(data_path)
    if len(base) + len(tests) > MAX_COLUMNS:
        raise ValueError("Excel 최대 16384열을 초과했습니다. 전체 측정값 Excel 출력을 해제하세요.")
    wb = Workbook(write_only=True)
    headers = list(base) + [f"{t['number']} · {t['title']}\n[{t['unit']}] {t['key']}" for t in tests]
    integer_fields = {'x_coord', 'y_coord', 'soft_bin', 'hard_bin', 'site'}
    float_fields = {'temp', 'test_time', 'plot_ul'}

    def new_sheet(page):
        ws = wb.create_sheet('전체 측정값' if page == 1 else f'전체 측정값 {page}')
        ws.freeze_panes = f'{get_column_letter(min(len(base), 5)+1)}2'
        ws.row_dimensions[1].height = 42
        for i in range(len(headers)):
            ws.column_dimensions[get_column_letter(i+1)].width = 24 if i >= len(base) else 16
        cells = []
        for title in headers:
            cell = WriteOnlyCell(ws, value=safe_cell(title))
            cell.font = Font(color='FFFFFF', bold=True)
            cell.alignment = Alignment(wrap_text=True, vertical='center')
            cell.fill = PatternFill('solid', fgColor='334155')
            cells.append(cell)
        ws.append(cells)
        return ws

    page, count, in_sheet = 1, 0, 0
    ws = new_sheet(page)
    try:
        for row in records(data_path):
            if row[0] != '1':
                continue
            if len(row) != offset + len(tests):
                raise ValueError('전체 측정값의 열 수가 올바르지 않습니다.')
            if in_sheet >= MAX_DATA_ROWS:
                ws.auto_filter.ref = f'A1:{get_column_letter(len(headers))}{in_sheet+1}'
                page += 1
                in_sheet = 0
                ws = new_sheet(page)
            values = []
            for name, value in zip(base, row[1:offset]):
                if name in integer_fields:
                    values.append(number(value, integer=True))
                elif name in float_fields:
                    values.append(number(value))
                else:
                    values.append(safe_cell(value))
            values.extend(number(v) for v in row[offset:])
            ws.append(values)
            count += 1
            in_sheet += 1
        ws.auto_filter.ref = f'A1:{get_column_letter(len(headers))}{in_sheet+1}'
        if count != summary['total']:
            raise ValueError('전체 측정값 Excel 행 수가 통계와 다릅니다.')
        definitions = wb.create_sheet('테스트 정의')
        definitions.freeze_panes = 'A2'
        definitions.append(['열 식별자', '테스트 번호', '테스트명', '단위', '하한', '상한'])
        for test in tests:
            definitions.append([safe_cell(test[k]) for k in ('key', 'number', 'title', 'unit', 'lower', 'upper')])
        info = wb.create_sheet('자료 안내')
        for label, value in [('입력 파일', Path(data_path).name), ('측정 행 수', count),
                             ('테스트 수', len(tests)), ('중간 데이터 포맷', summary['data_format']),
                             ('입력 SHA-256', summary['data_sha256']), ('수율 기준', summary['yield_policy']),
                             ('측정 시트 수', page), ('누락값', '빈 셀; 원본 STDF와 중간 데이터에서 추적')]:
            info.append([label, safe_cell(value)])
        wb.save(path)
    finally:
        wb.close()
