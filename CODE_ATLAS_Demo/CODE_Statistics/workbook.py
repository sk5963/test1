"""고정 포맷 통계 Excel 출력."""

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


def save_workbook(path, summary, tests, cells):
    wb = Workbook()
    ws = wb.active
    ws.title = "수율 요약"
    for k, label in [
        ("total", "전체 측정 행"),
        ("passed", "양품"),
        ("failed", "불량"),
        ("yield_percent", "수율(%)"),
        ("test_count", "테스트 수"),
    ]:
        ws.append([label, summary[k]])
    ws.append(["수율 기준", summary["yield_policy"]])
    ws = wb.create_sheet("테스트 통계")
    ws.append(
        [
            "열 식별자",
            "테스트 번호",
            "테스트명",
            "단위",
            "하한",
            "상한",
            "측정 수",
            "누락 수",
            "평균",
            "표본 표준편차",
            "최솟값",
            "최댓값",
            "한계 이탈 수",
        ]
    )
    for t in tests:
        ws.append(
            [
                safe_cell(t[k])
                for k in [
                    "key",
                    "number",
                    "title",
                    "unit",
                    "lower",
                    "upper",
                    "count",
                    "missing",
                    "mean",
                    "std",
                    "minimum",
                    "maximum",
                    "out_of_limit",
                ]
            ]
        )
    ws = wb.create_sheet("BIN 집계")
    ws.append(["구분", "BIN", "측정 수"])
    for kind in ("sbin", "hbin"):
        for b, n in summary[kind + "_counts"].items():
            ws.append([kind.upper(), safe_cell(b), n])
    ws = wb.create_sheet("맵 데이터")
    ws.append(["측정 행", "X", "Y", "양품", "사이트", "SBIN"])
    for c in cells:
        ws.append(
            [
                c["row"],
                c["x"],
                c["y"],
                c["passed"],
                safe_cell(c["site"]),
                safe_cell(c["sbin"]),
            ]
        )
    for ws in wb:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for c in ws[1]:
            c.font = Font(color="FFFFFF", bold=True)
            c.fill = PatternFill("solid", fgColor="334155")
        for column in ws.columns:
            letter = column[0].column_letter
            ws.column_dimensions[letter].width = min(
                48, max(14, max(len(str(c.value or "")) for c in column) + 2)
            )
    wb.save(path)


def safe_cell(value):
    return (
        "'" + value
        if isinstance(value, str) and value.startswith(("=", "+", "-", "@"))
        else value
    )
