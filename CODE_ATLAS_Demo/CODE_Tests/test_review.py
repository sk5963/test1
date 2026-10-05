"""전체 리뷰에서 찾은 장애·무결성·포맷 계약의 회귀 검증."""

import json
from pathlib import Path
import subprocess
import sys
import time
import unittest
from unittest.mock import patch
import numpy as np
import pyarrow.parquet as pq
from CODE_Tests import test_pipeline as baseline
from CODE_Common.io import ROOT, digest, read_json, write_json, file_lock, load_config
from CODE_Common.measurements import records, header, write_parquet
from CODE_Controller.pipeline import pipeline, create_run, execute_stage
from CODE_Controller.cleanup import cleanup
from CODE_STDF_Converter.service import convert
from CODE_Statistics.service import generate as statistics
from CODE_FTP.service import receive


class ReviewTests(unittest.TestCase):
    setUp = baseline.PipelineTests.setUp
    tearDown = baseline.PipelineTests.tearDown
    run_pipeline = baseline.PipelineTests.run_pipeline

    def cli(self, command, *args, input=None, module='CODE_Controller.cli'):
        return subprocess.run([sys.executable, '-m', module, command, '--config', str(self.config), *args],
                              cwd=ROOT, capture_output=True, text=True, input=input, timeout=15)

    def test_missing_completed_result_creates_new_run(self):
        old = self.run_pipeline()
        (self.store.root / 'runs' / old / 'statistics/map.svg').unlink()
        new = self.run_pipeline()
        self.assertNotEqual(old, new)
        self.assertEqual(self.store.run(old)['status'], '실패')
        self.assertTrue((self.store.root / 'runs' / new / 'statistics/map.svg').exists())

    def test_tampered_predecessor_blocks_report(self):
        run = create_run(self.store, self.source, self.c)
        for stage in ('convert', 'statistics'):
            execute_stage(self.store, self.c, run, stage, self.emit)
        summary = self.store.root / 'runs' / run / 'statistics/summary.json'
        content = read_json(summary); content['passed'] += 1; write_json(summary, content)
        with self.assertRaisesRegex(ValueError, '손상'):
            execute_stage(self.store, self.c, run, 'report', self.emit)
        self.assertFalse((summary.parent.parent / 'report').exists())

    def test_cleanup_preserves_data_without_original(self):
        run = self.run_pipeline()
        Path(self.store.run(run)['source']).unlink()
        result = cleanup(self.store, 0)
        self.assertEqual(result['removed_runs'], [])
        self.assertEqual(len(result['deferred_runs']), 1)
        self.assertTrue(Path(read_json(self.store.root / 'runs' / run / 'conversion/conversion.json')['data_path']).exists())

    def test_bad_cleanup_record_does_not_block_healthy_run(self):
        old = self.run_pipeline()
        (self.store.root / 'runs' / old / 'complete.json').write_text('{')
        self.c['histogram_bins'] += 1
        good = self.run_pipeline()
        result = cleanup(self.store, 0)
        self.assertEqual(result['removed_runs'], [good])
        self.assertEqual(result['deferred_runs'][0]['run_id'], old)

    def test_invalid_config_status_save_and_stop(self):
        self.config.write_text('{')
        status = self.cli('status')
        self.assertEqual(status.returncode, 0, status.stderr)
        self.assertTrue(json.loads(status.stdout)['config_error'])
        stop = self.cli('stop')
        self.assertEqual(stop.returncode, 0, stop.stderr)
        valid = {k: v for k, v in self.c.items() if not k.startswith('_')}
        saved = self.cli('save-config', input=json.dumps(valid))
        self.assertEqual(saved.returncode, 0, saved.stderr)
        self.assertEqual(load_config(self.config)['intermediate_format'], 'csv')
        valid['ftp']['tls'] = 'false'
        bad = self.cli('save-config', input=json.dumps(valid))
        self.assertEqual(bad.returncode, 1)
        self.assertEqual(json.loads(bad.stdout)['type'], 'error')
        self.assertFalse(load_config(self.config)['ftp']['tls'])

    def test_standalone_ftp_uses_shared_lock_and_upload_history(self):
        with file_lock(self.store.root / 'pipeline.lock'):
            result = self.cli('download', module='CODE_FTP')
        self.assertEqual(result.returncode, 1)
        self.assertIn('다른 작업', result.stdout)
        run = create_run(self.store, self.source, self.c)
        for stage in ('convert', 'statistics', 'report'):
            execute_stage(self.store, self.c, run, stage, self.emit)
        uploaded = self.cli('upload', '--run-id', run, module='CODE_FTP')
        self.assertEqual(uploaded.returncode, 0, uploaded.stdout + uploaded.stderr)
        self.assertEqual(self.store.run(run)['status'], '완료')
        self.assertEqual(self.store.steps(run)['upload']['status'], '성공')

    def test_different_ftp_origins_do_not_overwrite_raw(self):
        a = receive(self.c, self.store)['files'][0]
        other = {**self.c, 'ftp': {**self.c['ftp'], 'remote_input': '/other/input'}}
        b = receive(other, self.store)['files'][0]
        self.assertNotEqual(a, b)
        self.assertEqual(digest(a), digest(b))
        self.assertEqual(self.store.db.execute('SELECT count(*) FROM receipts').fetchone()[0], 2)

    def test_changed_csv_between_statistical_passes_is_rejected(self):
        converted = convert(self.source, Path(self.temp.name) / 'conversion', data_format='csv')
        csv = Path(converted['data_path'])
        def progress(n, message):
            if n == 45:
                csv.write_text(csv.read_text().replace('1.25', '1.26', 1))
        with self.assertRaisesRegex(ValueError, '변경'):
            statistics(csv, Path(self.temp.name) / 'statistics', progress=progress)
        self.assertFalse((Path(self.temp.name) / 'statistics/summary.json').exists())

    def test_parquet_matches_csv_and_preserves_definition(self):
        a = convert(self.source, Path(self.temp.name) / 'csv_conversion', data_format='csv')
        b = convert(self.source, Path(self.temp.name) / 'parquet_conversion', data_format='parquet')
        self.assertIsNone(b['full_csv'])
        self.assertFalse(list(Path(b['data_path']).parent.glob('*.csv')))
        self.assertEqual(header(a['data_path']), header(b['data_path']))
        csv_rows = [r for r in records(a['data_path']) if r[0] == '1']
        pq_rows = [r for r in records(b['data_path']) if r[0] == '1']
        offset = header(a['data_path'])[2]
        for left, right in zip(csv_rows, pq_rows):
            self.assertEqual(left[1:offset], right[1:offset])
            np.testing.assert_array_equal(np.array(left[offset:], dtype=float), np.array(right[offset:], dtype=float))
        self.assertEqual(len(csv_rows), len(pq_rows))
        with pq.ParquetFile(b['data_path']) as f:
            self.assertEqual(f.metadata.row_group(0).column(offset-1).compression, 'ZSTD')
            self.assertEqual(f.metadata.num_rows, 24)
            source_meta = json.loads(f.metadata.metadata[b'atlas.source_metadata'])
            self.assertEqual(source_meta['source_sha256'], digest(self.source))
        s1 = statistics(a['data_path'], Path(self.temp.name) / 'stats_csv')
        s2 = statistics(b['data_path'], Path(self.temp.name) / 'stats_parquet')
        for key in ('total', 'passed', 'failed', 'yield_percent', 'test_count', 'sbin_counts', 'hbin_counts'):
            self.assertEqual(s1[key], s2[key])
        self.assertEqual(read_json(Path(self.temp.name) / 'stats_csv/histograms.json'),
                         read_json(Path(self.temp.name) / 'stats_parquet/histograms.json'))
        self.assertEqual(read_json(Path(self.temp.name) / 'stats_csv/map.json'),
                         read_json(Path(self.temp.name) / 'stats_parquet/map.json'))

    def test_parquet_pipeline_cleanup_and_result_recovery(self):
        self.c['intermediate_format'] = 'parquet'
        run = self.run_pipeline()
        folder = self.store.root / 'runs' / run
        data = Path(read_json(folder / 'conversion/conversion.json')['data_path'])
        self.assertTrue(data.exists())
        manifest = read_json(folder / 'complete.json')
        self.assertFalse(any(f['path'].endswith(('.csv', '.parquet')) for f in manifest['files']))
        self.assertEqual(cleanup(self.store, 0)['removed_runs'], [run])
        self.assertFalse(data.exists())
        self.assertEqual(pipeline(self.store, self.c, self.emit)['completed'], [])
        (folder / 'report/index.html').write_text('bad')
        new = self.run_pipeline()
        self.assertNotEqual(new, run)

    def test_parquet_nulls_duplicate_numbers_and_multiple_row_groups(self):
        a = convert(self.source, Path(self.temp.name) / 'conversion', data_format='csv')
        import csv
        path = Path(a['data_path'])
        with path.open(newline='', encoding='utf-8-sig') as f:
            rows = list(csv.reader(f))
        offset = header(path)[2]
        for row in rows:
            if row and row[0] == '0.2':
                row[offset+1] = row[offset]
        measurement = next(r for r in rows if r and r[0] == '1')
        measurement[offset] = ''
        measurement[offset+1] = 'nan'
        with path.open('w', newline='', encoding='utf-8') as f:
            csv.writer(f).writerows(rows)
        target = path.with_suffix('.parquet')
        write_parquet(path, target, row_group_size=5)
        with pq.ParquetFile(target) as f:
            self.assertEqual(f.metadata.num_row_groups, 5)
        _, tests, _ = header(target)
        self.assertEqual(tests[0]['number'], tests[1]['number'])
        self.assertNotEqual(tests[0]['key'], tests[1]['key'])
        result = statistics(target, Path(self.temp.name) / 'statistics')
        hist = read_json(Path(self.temp.name) / 'statistics/histograms.json')['tests']
        self.assertEqual(result['test_count'], 20)
        self.assertEqual((hist[0]['count'], hist[0]['missing']), (23, 1))
        self.assertEqual((hist[1]['count'], hist[1]['missing']), (23, 1))

    def test_failed_parquet_export_keeps_bridge_csv(self):
        destination = Path(self.temp.name) / 'conversion'
        with patch('CODE_STDF_Converter.service.write_parquet', side_effect=ValueError('저장 실패')):
            with self.assertRaises(ValueError):
                convert(self.source, destination, data_format='parquet')
        self.assertTrue(list(destination.glob('*.csv')))
        self.assertFalse((destination / 'conversion.json').exists())

    def test_measurement_excel_values_definitions_and_upload(self):
        from openpyxl import load_workbook
        self.c['intermediate_format'] = 'parquet'
        run = self.run_pipeline()
        folder = self.store.root / 'runs' / run
        meta = read_json(folder / 'conversion/conversion.json')
        _, tests, offset = header(meta['data_path'])
        expected = [r for r in records(meta['data_path']) if r[0] == '1']
        wb = load_workbook(folder / 'statistics/measurements.xlsx', read_only=True)
        try:
            self.assertEqual(wb.sheetnames, ['전체 측정값', '테스트 정의', '자료 안내'])
            rows = list(wb['전체 측정값'].values)
            self.assertEqual(len(rows), 25)
            self.assertEqual(len(rows[0]), offset - 1 + len(tests))
            self.assertIn(tests[0]['number'], rows[0][offset-1])
            for actual, source in zip(rows[1:], expected):
                np.testing.assert_allclose(actual[offset-1:], source[offset:], rtol=1e-15)
            self.assertEqual(len(list(wb['테스트 정의'].values)), 21)
        finally:
            wb.close()
        published = self.store.root / 'demo_remote/results' / run / 'statistics/measurements.xlsx'
        self.assertEqual(digest(published), digest(folder / 'statistics/measurements.xlsx'))
        self.assertEqual(cleanup(self.store, 0)['removed_runs'], [run])
        self.assertTrue(published.exists())
        self.assertTrue((folder / 'statistics/measurements.xlsx').exists())

    def test_measurement_excel_splits_rows_and_can_be_disabled(self):
        from openpyxl import load_workbook
        from CODE_Statistics.measurement_workbook import save_measurements
        converted = convert(self.source, Path(self.temp.name) / 'conversion', data_format='parquet')
        output = Path(self.temp.name) / 'statistics'
        summary = statistics(converted['data_path'], output, measurement_excel=False)
        self.assertFalse((output / 'measurements.xlsx').exists())
        target = Path(self.temp.name) / 'split.xlsx'
        with patch('CODE_Statistics.measurement_workbook.MAX_DATA_ROWS', 10):
            save_measurements(target, converted['data_path'], summary)
        wb = load_workbook(target, read_only=True)
        try:
            self.assertEqual(wb.sheetnames[:3], ['전체 측정값', '전체 측정값 2', '전체 측정값 3'])
            self.assertEqual([len(list(wb[n].values))-1 for n in wb.sheetnames[:3]], [10, 10, 4])
        finally:
            wb.close()

    def test_output_format_change_creates_separate_run(self):
        csv_run = self.run_pipeline()
        csv_data = Path(read_json(self.store.root / 'runs' / csv_run / 'conversion/conversion.json')['data_path'])
        self.c['intermediate_format'] = 'parquet'
        parquet_run = self.run_pipeline()
        self.assertNotEqual(csv_run, parquet_run)
        self.assertTrue(csv_data.exists())
        self.assertEqual(read_json(self.store.root / 'runs' / parquet_run / 'conversion/conversion.json')['data_format'], 'parquet')

    def test_watch_survives_invalid_config_and_can_stop(self):
        config = {k: v for k, v in self.c.items() if not k.startswith('_')}
        config['poll_seconds'] = 5
        write_json(self.config, config)
        log = Path(self.temp.name) / 'watch.log'
        with log.open('w') as out:
            proc = subprocess.Popen([sys.executable, '-m', 'CODE_Controller.cli', 'watch', '--config', str(self.config)],
                                    cwd=ROOT, stdout=out, stderr=out)
            try:
                deadline = time.monotonic()+15
                while not self.store.db.execute("SELECT count(*) FROM runs WHERE status='완료'").fetchone()[0]:
                    if time.monotonic() > deadline or proc.poll() is not None:
                        self.fail(log.read_text())
                    time.sleep(.1)
                self.config.write_text('{')
                while '자동 처리 오류' not in log.read_text():
                    if time.monotonic() > deadline or proc.poll() is not None:
                        self.fail(log.read_text())
                    time.sleep(.1)
                self.assertIsNone(proc.poll())
                self.assertEqual(self.cli('stop').returncode, 0)
                self.assertEqual(proc.wait(timeout=5), 0)
            finally:
                if proc.poll() is None:
                    proc.terminate(); proc.wait(timeout=5)


if __name__ == '__main__':
    unittest.main()
