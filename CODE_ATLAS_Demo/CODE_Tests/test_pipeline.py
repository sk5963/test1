"""실제 합성 STDF로 단계 계약·재실행·정리 및 서버 제어를 검증한다."""

import json
from pathlib import Path
import subprocess
import sys
import tempfile
import time
import unittest
import numpy as np
from openpyxl import load_workbook
from CODE_Common.io import ROOT, read_json, write_json, load_config, digest, file_lock
from CODE_Common.store import Store
from CODE_Common.full_csv import records, header
from CODE_Controller.pipeline import pipeline, create_run, execute_stage
from CODE_Controller.cleanup import cleanup
from CODE_Demo_Data.generate import generate
from CODE_FTP.service import receive


class PipelineTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.config = Path(self.temp.name) / "config.json"
        c = read_json(ROOT / "config.json")
        c["transport"] = "demo"
        c["intermediate_format"] = "csv"
        write_json(self.config, c)
        self.c = load_config(self.config)
        self.store = Store(self.c["_workspace"])
        self.source = generate(
            self.store.root / "demo_remote/incoming", parts=24, tests=20
        )
        self.emit = lambda *a: None

    def tearDown(self):
        self.store.close()
        self.temp.cleanup()

    def run_pipeline(self):
        result = pipeline(self.store, self.c, self.emit)
        self.assertEqual(result["failed"], [])
        return result["completed"][0]

    def test_end_to_end_and_reference_statistics(self):
        run = self.run_pipeline()
        folder = self.store.root / "runs" / run
        summary = read_json(folder / "statistics/summary.json")
        expected = read_json(self.source.parent / "expected.json")
        self.assertEqual(
            (summary["total"], summary["passed"], summary["failed"]),
            (24, expected["pass"], expected["fail"]),
        )
        conversion = read_json(folder / "conversion/conversion.json")
        _, _, offset = header(conversion["full_csv"])
        values = np.array(
            [float(r[offset]) for r in records(conversion["full_csv"]) if r[0] == "1"]
        )
        first = read_json(folder / "statistics/histograms.json")["tests"][0]
        self.assertAlmostEqual(first["mean"], float(values.mean()), places=12)
        self.assertAlmostEqual(first["std"], float(values.std(ddof=1)), places=12)
        self.assertEqual(sum(first["counts"]), 24)
        wb = load_workbook(folder / "statistics/statistics.xlsx", read_only=True)
        self.assertEqual(
            wb.sheetnames, ["수율 요약", "테스트 통계", "BIN 집계", "맵 데이터"]
        )
        wb.close()
        remote = self.store.root / "demo_remote/results" / run
        manifest = read_json(remote / "complete.json")
        self.assertFalse(any(f["path"].endswith(".csv") for f in manifest["files"]))
        for f in manifest["files"]:
            self.assertEqual(digest(remote / f["path"]), f["sha256"])

    def test_duplicate_skip_and_local_corruption_recovery(self):
        self.run_pipeline()
        again = pipeline(self.store, self.c, self.emit)
        self.assertEqual(again["receipt"]["skipped"], 1)
        self.assertEqual(again["completed"], [])
        raw = Path(self.store.db.execute("SELECT local_path FROM receipts").fetchone()[0])
        raw.write_bytes(b"bad")
        result = receive(self.c, self.store)
        self.assertEqual(result["downloaded"], 1)
        self.assertEqual(digest(raw), digest(self.source))
        # 운영 주기 변경은 불필요한 전체 재처리를 발생시키지 않는다.
        self.c["poll_seconds"] = 10
        self.assertEqual(pipeline(self.store, self.c, self.emit)["completed"], [])

    def test_manual_stages_and_input_integrity(self):
        run = create_run(self.store, self.source, self.c)
        with self.assertRaises(ValueError):
            execute_stage(self.store, self.c, run, "statistics", self.emit)
        execute_stage(self.store, self.c, run, "convert", self.emit)
        csv_path = Path(
            json.loads(self.store.steps(run)["convert"]["result_json"])["full_csv"]
        )
        old = csv_path.read_bytes()
        csv_path.write_bytes(b"bad")
        with self.assertRaises(ValueError):
            execute_stage(self.store, self.c, run, "statistics", self.emit)
        self.assertEqual(self.store.run(run)["status"], "실패")
        csv_path.write_bytes(old)
        for stage in ("statistics", "report", "upload"):
            execute_stage(self.store, self.c, run, stage, self.emit)
        self.assertEqual(self.store.run(run)["status"], "완료")

    def test_failed_conversion_is_not_cleaned(self):
        self.source.write_bytes(self.source.read_bytes()[:-25])
        result = pipeline(self.store, self.c, self.emit)
        self.assertEqual(len(result["failed"]), 1)
        self.assertEqual(cleanup(self.store, 0)["removed_runs"], [])
        self.assertEqual(self.store.runs()[0]["status"], "실패")

    def test_cleanup_only_verified_success_csv(self):
        run = self.run_pipeline()
        folder = self.store.root / "runs" / run
        conversion = read_json(folder / "conversion/conversion.json")
        csv_path = Path(conversion["full_csv"])
        report = folder / "report/index.html"
        original = report.read_bytes()
        report.write_bytes(b"corrupt")
        self.assertEqual(cleanup(self.store, 0)["removed_runs"], [])
        self.assertTrue(csv_path.exists())
        report.write_bytes(original)
        self.assertEqual(cleanup(self.store, 0)["removed_runs"], [run])
        self.assertFalse(csv_path.exists())
        self.assertTrue(self.source.exists())
        self.assertTrue(Path(self.store.run(run)["source"]).exists())
        self.assertTrue(report.exists())
        self.assertEqual(pipeline(self.store, self.c, self.emit)["completed"], [])

    def test_shared_lock_rejects_overlap(self):
        with file_lock(self.store.root / "pipeline.lock"):
            with self.assertRaises(RuntimeError):
                with file_lock(self.store.root / "pipeline.lock"):
                    pass

    def test_vendor_bytes_unchanged(self):
        manifest = read_json(ROOT / "CODE_Vendor/SOURCE_MANIFEST.json")
        for f in manifest["files"]:
            self.assertEqual(digest(ROOT / f["path"]), f["sha256"])

    def test_watch_runs_and_cooperative_stop(self):
        log = Path(self.temp.name) / "watch.log"
        with log.open("w", encoding="utf-8") as f:
            process = subprocess.Popen(
                [
                    sys.executable,
                    "-m",
                    "CODE_Controller.cli",
                    "watch",
                    "--config",
                    str(self.config),
                ],
                cwd=ROOT,
                stdout=f,
                stderr=f,
            )
            try:
                deadline = time.monotonic() + 15
                while time.monotonic() < deadline:
                    if self.store.db.execute(
                        "SELECT count(*) FROM runs WHERE status='완료'"
                    ).fetchone()[0]:
                        break
                    if process.poll() is not None:
                        self.fail(log.read_text())
                    time.sleep(0.1)
                else:
                    self.fail("서버 처리 제한시간 초과")
                request = subprocess.run(
                    [
                        sys.executable,
                        "-m",
                        "CODE_Controller.cli",
                        "stop",
                        "--config",
                        str(self.config),
                    ],
                    cwd=ROOT,
                    capture_output=True,
                    text=True,
                    timeout=5,
                )
                self.assertEqual(request.returncode, 0, request.stderr)
                self.assertEqual(process.wait(timeout=5), 0)
                self.assertFalse(read_json(self.store.root / "service.json")["running"])
            finally:
                if process.poll() is None:
                    process.terminate()
                    process.wait(timeout=5)


if __name__ == "__main__":
    unittest.main()
